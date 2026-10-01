import io
from docx import Document
from docx.shared import Pt, RGBColor, Inches
from docx.enum.section import WD_ORIENT
from openpyxl import Workbook
from openpyxl.styles import Font, PatternFill, Alignment
from openpyxl.utils import get_column_letter
from reportlab.lib.pagesizes import A4, landscape
from reportlab.lib import colors
from reportlab.lib.styles import getSampleStyleSheet, ParagraphStyle
from reportlab.platypus import SimpleDocTemplate, Paragraph, Table, TableStyle, Spacer
from reportlab.pdfgen import canvas as rl_canvas
from reportlab.lib.utils import simpleSplit
from pptx import Presentation
from pptx.util import Inches as PI, Pt as PPt
from pptx.dml.color import RGBColor as PRGB
from pptx.chart.data import CategoryChartData
from pptx.enum.chart import XL_CHART_TYPE, XL_LEGEND_POSITION
from pathlib import Path
from xml.sax.saxutils import escape as xesc

GREEN = "1A4D3A"
NA = "Belum tersedia"
LOGO = str(Path(__file__).parent / "logo.png")
STATUS_TXT = {"tercapai": "Tercapai", "belum_tercapai": "Belum tercapai", "belum_tersedia": NA,
              "sasaran_kosong": "Sasaran belum diisi", "target_kosong": "Target belum ditetapkan"}
METHOD_TXT = {"kumulatif": "Kumulatif", "rasio": "Rasio periode", "terakhir": "Nilai terakhir"}
REPORT_STATUS = {"draf": "Draf", "diajukan": "Diajukan", "perlu_perbaikan": "Perlu perbaikan", "disetujui": "Disetujui"}
HEAD = ["Indikator", "Satuan", "Metode", "Sasaran", "Target periode (%)", "Capaian", "Capaian (%)", "Status"]
NARR = [("kendala", "Kendala"), ("upaya", "Upaya yang sudah dilakukan"), ("hasil_upaya", "Hasil upaya"),
        ("rtl", "Rencana tindak lanjut"), ("dukungan", "Dukungan yang dibutuhkan")]


def fmt(v, pct=False):
    if v is None:
        return NA
    s = f"{v:,.0f}" if float(v).is_integer() else f"{v:,.2f}"
    s = s.replace(",", "X").replace(".", ",").replace("X", ".")
    return s + ("%" if pct else "")


def ind_row(i):
    r = i["recap"]
    return [i["name"], i.get("unit") or "-", METHOD_TXT.get(i.get("method"), "-"), fmt(r["sasaran"]),
            fmt(r["target_periode"]), fmt(r["capaian"]), fmt(r["persen"], True), STATUS_TXT[r["status"]]]


def title_of(ds):
    u = {"spm": "SPM", "program": "Program"}.get(ds.get("kind"), "SPM & Program")
    if len(ds["programs"]) == 1:
        p = ds["programs"][0]
        return f"Laporan {'SPM' if p.get('kind') == 'spm' else 'Program'} {p['name']}"
    return f"Laporan Kolektif {u}"


def report_docx(ds):
    d = Document()
    sec = d.sections[0]
    sec.orientation = WD_ORIENT.LANDSCAPE
    sec.page_width, sec.page_height = sec.page_height, sec.page_width
    d.add_picture(LOGO, width=Inches(2.2))
    h = d.add_heading(title_of(ds), 0)
    h.runs[0].font.color.rgb = RGBColor.from_string(GREEN)
    d.add_paragraph(f"UPT Puskesmas Melati - Periode: {ds['periode']}")
    for p in ds["programs"]:
        d.add_heading(p["name"], 1)
        d.add_paragraph(f"Penanggung jawab: {p.get('pj_name') or '-'}")
        t = d.add_table(rows=1, cols=len(HEAD))
        t.style = "Light Grid Accent 6"
        for c, txt in zip(t.rows[0].cells, HEAD):
            c.text = txt
        for i in p["indicators"]:
            for c, txt in zip(t.add_row().cells, ind_row(i)):
                c.text = txt
        if not p["indicators"]:
            d.add_paragraph("Indikator belum ditetapkan.")
        d.add_heading("Narasi laporan", 2)
        if not p["narasi"]:
            d.add_paragraph(f"{NA} - belum ada laporan yang diajukan pada periode ini.")
        for n in p["narasi"]:
            d.add_paragraph(f"{n['bulan']} ({REPORT_STATUS.get(n['status'], n['status'])})").runs[0].bold = True
            for k, lbl in NARR:
                para = d.add_paragraph(style="List Bullet")
                para.add_run(f"{lbl}: ").bold = True
                para.add_run(n.get(k) or NA)
    d.add_paragraph().add_run("Catatan: nilai 'Belum tersedia' berarti data belum dilaporkan, bukan nol. "
                              "Rekap triwulan/tahunan dihitung dari laporan bulanan sesuai metode tiap indikator.").italic = True
    buf = io.BytesIO()
    d.save(buf)
    return buf.getvalue()


def report_xlsx(ds):
    wb = Workbook()
    ws = wb.active
    ws.title = "Rekap"
    fill = PatternFill("solid", fgColor=GREEN)
    ws.append([title_of(ds)])
    ws.append([f"UPT Puskesmas Melati - Periode: {ds['periode']}"])
    ws.append([])
    ws.append(["SPM/Program", "PJ"] + HEAD)
    for c in ws[4]:
        c.font, c.fill = Font(bold=True, color="FFFFFF"), fill
    for p in ds["programs"]:
        for i in p["indicators"]:
            ws.append([p["name"], p.get("pj_name") or "-"] + ind_row(i))
    for idx, w in enumerate([22, 18, 40, 12, 14, 12, 14, 12, 12, 20], 1):
        ws.column_dimensions[get_column_letter(idx)].width = w
    ws2 = wb.create_sheet("Narasi")
    ws2.append(["SPM/Program", "Bulan", "Status"] + [l for _, l in NARR])
    for c in ws2[1]:
        c.font, c.fill = Font(bold=True, color="FFFFFF"), fill
    for p in ds["programs"]:
        if not p["narasi"]:
            ws2.append([p["name"], ds["periode"], NA] + [NA] * len(NARR))
        for n in p["narasi"]:
            ws2.append([p["name"], n["bulan"], REPORT_STATUS.get(n["status"])] + [n.get(k) or NA for k, _ in NARR])
    for idx in range(1, 9):
        ws2.column_dimensions[get_column_letter(idx)].width = 30
        for row in ws2.iter_rows(min_row=2):
            row[idx - 1].alignment = Alignment(wrap_text=True, vertical="top")
    buf = io.BytesIO()
    wb.save(buf)
    return buf.getvalue()


def _styles():
    ss = getSampleStyleSheet()
    cell = ParagraphStyle("cell", parent=ss["BodyText"], fontSize=8, leading=10)
    return ss, cell


def report_pdf(ds):
    buf = io.BytesIO()
    doc = SimpleDocTemplate(buf, pagesize=landscape(A4), leftMargin=30, rightMargin=30, topMargin=30, bottomMargin=30)
    ss, cell = _styles()
    h = ParagraphStyle("h", parent=ss["Title"], textColor=colors.HexColor("#" + GREEN))
    el = [Paragraph(xesc(title_of(ds)), h), Paragraph(f"UPT Puskesmas Melati - Periode: {xesc(ds['periode'])}", ss["Normal"]), Spacer(1, 10)]
    for p in ds["programs"]:
        el.append(Paragraph(xesc(p["name"]), ss["Heading2"]))
        el.append(Paragraph(f"Penanggung jawab: {xesc(p.get('pj_name') or '-')}", ss["Normal"]))
        data = [[Paragraph(f"<b>{x}</b>", cell) for x in HEAD]] + [[Paragraph(xesc(x), cell) for x in ind_row(i)] for i in p["indicators"]]
        if len(data) > 1:
            t = Table(data, colWidths=[200, 60, 70, 60, 70, 60, 60, 90], repeatRows=1)
            t.setStyle(TableStyle([("BACKGROUND", (0, 0), (-1, 0), colors.HexColor("#E8F5E9")),
                                   ("GRID", (0, 0), (-1, -1), 0.4, colors.HexColor("#9CB8A8")),
                                   ("VALIGN", (0, 0), (-1, -1), "TOP")]))
            el.append(t)
        el.append(Spacer(1, 6))
        if not p["narasi"]:
            el.append(Paragraph(f"<i>Narasi: {NA}</i>", ss["Normal"]))
        for n in p["narasi"]:
            el.append(Paragraph(f"<b>{n['bulan']}</b> ({REPORT_STATUS.get(n['status'])})", ss["Normal"]))
            for k, lbl in NARR:
                el.append(Paragraph(f"<b>{lbl}:</b> {xesc(n.get(k) or NA)}", cell))
        el.append(Spacer(1, 12))
    el.append(Paragraph("<i>'Belum tersedia' berarti data belum dilaporkan, bukan nol.</i>", cell))
    doc.build(el)
    return buf.getvalue()


def _bullets(s):
    return [b for b in (s.get("bullets") or []) if str(b).strip()]


def pres_pptx(pres):
    prs = Presentation()
    prs.slide_width, prs.slide_height = PI(13.333), PI(7.5)
    blank = prs.slide_layouts[6]
    green = PRGB.from_string(GREEN)
    for s in pres["slides"]:
        sl = prs.slides.add_slide(blank)
        if s.get("layout") == "title":
            bg = sl.background.fill
            bg.solid()
            bg.fore_color.rgb = PRGB.from_string("F8FAF7")
            sl.shapes.add_picture(LOGO, PI(0.8), PI(0.6), height=PI(1.4))
            tb = sl.shapes.add_textbox(PI(0.8), PI(2.8), PI(11.5), PI(2)).text_frame
            tb.word_wrap = True
            tb.text = s.get("title", "")
            tb.paragraphs[0].runs[0].font.size, tb.paragraphs[0].runs[0].font.bold = PPt(40), True
            tb.paragraphs[0].runs[0].font.color.rgb = green
            p = tb.add_paragraph()
            p.text = s.get("subtitle", "")
            p.runs[0].font.size = PPt(20) if p.runs else PPt(20)
            continue
        bar = sl.shapes.add_shape(1, 0, 0, prs.slide_width, PI(1.1))
        bar.fill.solid()
        bar.fill.fore_color.rgb = green
        bar.line.fill.background()
        tf = bar.text_frame
        tf.text = s.get("title", "")
        tf.paragraphs[0].runs[0].font.size, tf.paragraphs[0].runs[0].font.bold = PPt(28), True
        layout = s.get("layout")
        if layout == "table" and s.get("table", {}).get("rows"):
            hdr, rows = s["table"]["headers"], s["table"]["rows"]
            shp = sl.shapes.add_table(len(rows) + 1, len(hdr), PI(0.5), PI(1.4), PI(12.3), PI(0.4) * (len(rows) + 1))
            for j, hname in enumerate(hdr):
                shp.table.cell(0, j).text = str(hname)
            for ri, row in enumerate(rows, 1):
                for j, v in enumerate(row):
                    shp.table.cell(ri, j).text = str(v)
                    shp.table.cell(ri, j).text_frame.paragraphs[0].runs[0].font.size = PPt(12)
        elif layout == "chart" and s.get("chart", {}).get("labels"):
            cd = CategoryChartData()
            cd.categories = s["chart"]["labels"]
            for se in s["chart"]["series"]:
                cd.add_series(se["name"], [v if v is not None else None for v in se["data"]])
            ch = sl.shapes.add_chart(XL_CHART_TYPE.COLUMN_CLUSTERED, PI(0.7), PI(1.4), PI(12), PI(5.6), cd).chart
            ch.has_legend, ch.legend.position, ch.legend.include_in_layout = True, XL_LEGEND_POSITION.BOTTOM, False
        else:
            tb = sl.shapes.add_textbox(PI(0.7), PI(1.4), PI(12), PI(5.6)).text_frame
            tb.word_wrap = True
            for k, b in enumerate(_bullets(s) or [NA]):
                p = tb.paragraphs[0] if k == 0 else tb.add_paragraph()
                p.text = f"\u2022 {b}"
                p.runs[0].font.size = PPt(18)
                p.space_after = PPt(8)
        if s.get("notes"):
            sl.notes_slide.notes_text_frame.text = s["notes"]
    buf = io.BytesIO()
    prs.save(buf)
    return buf.getvalue()


def pres_pdf(pres):
    W, H = 960, 540
    buf = io.BytesIO()
    c = rl_canvas.Canvas(buf, pagesize=(W, H))
    green = colors.HexColor("#" + GREEN)
    _, cell = _styles()
    for s in pres["slides"]:
        c.setFillColor(colors.HexColor("#F8FAF7"))
        c.rect(0, 0, W, H, fill=1, stroke=0)
        if s.get("layout") == "title":
            c.drawImage(LOGO, 60, H - 160, width=240, height=120, preserveAspectRatio=True, mask="auto")
            c.setFillColor(green)
            c.setFont("Helvetica-Bold", 34)
            y = H - 260
            for line in simpleSplit(s.get("title", ""), "Helvetica-Bold", 34, W - 120):
                c.drawString(60, y, line)
                y -= 42
            c.setFillColor(colors.HexColor("#334155"))
            c.setFont("Helvetica", 18)
            c.drawString(60, y - 10, s.get("subtitle", ""))
            c.showPage()
            continue
        c.setFillColor(green)
        c.rect(0, H - 80, W, 80, fill=1, stroke=0)
        c.setFillColor(colors.white)
        c.setFont("Helvetica-Bold", 24)
        c.drawString(40, H - 52, s.get("title", "")[:80])
        layout = s.get("layout")
        if layout == "table" and s.get("table", {}).get("rows"):
            data = [[Paragraph(f"<b>{xesc(str(x))}</b>", cell) for x in s["table"]["headers"]]] + \
                   [[Paragraph(xesc(str(v)), cell) for v in r] for r in s["table"]["rows"]]
            t = Table(data, colWidths=[(W - 80) / len(data[0])] * len(data[0]))
            t.setStyle(TableStyle([("BACKGROUND", (0, 0), (-1, 0), colors.HexColor("#E8F5E9")),
                                   ("GRID", (0, 0), (-1, -1), 0.4, colors.HexColor("#9CB8A8"))]))
            _, th = t.wrapOn(c, W - 80, H - 120)
            t.drawOn(c, 40, H - 100 - th)
        elif layout == "chart" and s.get("chart", {}).get("labels"):
            labels, series = s["chart"]["labels"], s["chart"]["series"]
            vals = [v for se in series for v in se["data"] if v is not None]
            mx = max(vals + [100])
            x0, y0, cw, chh = 70, 90, W - 120, H - 220
            c.setStrokeColor(colors.HexColor("#94A3B8"))
            c.line(x0, y0, x0 + cw, y0)
            gw = cw / max(len(labels), 1)
            pal = [green, colors.HexColor("#F59E0B"), colors.HexColor("#10B981")]
            bw = gw * 0.7 / max(len(series), 1)
            for gi, lab in enumerate(labels):
                for si, se in enumerate(series):
                    v = se["data"][gi]
                    bx = x0 + gi * gw + gw * 0.15 + si * bw
                    c.setFillColor(pal[si % 3])
                    if v is None:
                        c.setFont("Helvetica-Oblique", 7)
                        c.setFillColor(colors.HexColor("#94A3B8"))
                        c.drawString(bx, y0 + 4, "N/A")
                    else:
                        bh = chh * v / mx
                        c.rect(bx, y0, bw - 2, bh, fill=1, stroke=0)
                        c.setFont("Helvetica", 8)
                        c.drawString(bx, y0 + bh + 3, fmt(v))
                c.setFillColor(colors.HexColor("#334155"))
                c.setFont("Helvetica", 8)
                for li, line in enumerate(simpleSplit(str(lab), "Helvetica", 8, gw - 6)[:3]):
                    c.drawString(x0 + gi * gw + 3, y0 - 12 - li * 10, line)
            for si, se in enumerate(series):
                c.setFillColor(pal[si % 3])
                c.rect(x0 + si * 160, 30, 10, 10, fill=1, stroke=0)
                c.setFillColor(colors.black)
                c.drawString(x0 + si * 160 + 14, 31, se["name"])
            c.setFont("Helvetica-Oblique", 8)
            c.drawString(x0, 15, "N/A = Belum tersedia")
        else:
            y = H - 120
            c.setFillColor(colors.HexColor("#1f2937"))
            for b in _bullets(s) or [NA]:
                for li, line in enumerate(simpleSplit(str(b), "Helvetica", 15, W - 110)):
                    c.setFont("Helvetica", 15)
                    c.drawString(60, y, ("\u2022 " if li == 0 else "   ") + line)
                    y -= 21
                y -= 8
                if y < 30:
                    break
        c.showPage()
    c.save()
    return buf.getvalue()
